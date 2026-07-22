"""Оформление по ГОСТ 7.32-2017 (выписка «Структура и правила оформления»)
и парность метрик с превью (frontend/src/lib/gostRender.ts):

- строка «полуторного» интервала Word = 24,15 пт (превью line-height:1.725);
- заголовки БЕЗ доп. интервалов до/после — везде обычный полуторный интервал;
- перечисления — с красной строки, продолжение к левому полю (без висячего);
- содержание: подразделы 0,5 см, пункты 1 см; реферат до содержания и не в нём;
- источники: номер С точкой, с абзацного отступа; заголовки без точки в конце;
- свободные строки вокруг блоков — НАСТОЯЩИЕ пустые абзацы (не space_before);
- подписи рисунков и таблиц — одинарный интервал, подпись таблицы вплотную к ней;
- заглушка картинки 125×62 мм (как в превью);
- номера страниц — кеглем основного текста; px→мм по 96 dpi без явной плотности.

Фикстура data/kalman-report.md — реальная курсовая (копия
frontend/src/fixtures/kalman-report.md; менять только парой).
"""

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.convert.gost import FREE_LINE, LINE_PT, build_docx
from app.convert.models import GostSettings

FIXTURE = (Path(__file__).parent / "data" / "kalman-report.md").read_text(encoding="utf-8")

NBSP = " "


def _settings(**kw) -> GostSettings:
    base = dict(title_page=False, toc=False, page_numbers=False)
    base.update(kw)
    return GostSettings(**base)


def _doc(md: str, settings: GostSettings | None = None, assets: dict | None = None) -> Document:
    return Document(io.BytesIO(build_docx(md, settings or _settings(), assets or {})))


# ---------- 2 «Общие требования» ----------


def test_line_metric_matches_word():
    # Свободная строка/строка текста = 14пт × 1,15 (hhea Times New Roman) × 1,5.
    # Ровно эта же величина в превью: line-height 1.725 (LINE_HEIGHT).
    assert LINE_PT == Pt(24.15)
    assert FREE_LINE == Pt(24.15)


def test_normal_style_gost():
    doc = _doc("Просто абзац текста.")
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(14)
    pf = normal.paragraph_format
    assert pf.line_spacing == 1.5
    assert pf.space_before == Pt(0) and pf.space_after == Pt(0)
    p = next(p for p in doc.paragraphs if "Просто абзац" in p.text)
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert round(p.paragraph_format.first_line_indent.mm, 1) == 12.5


def test_page_number_same_font_size_as_text():
    # «Номера страниц — тем же шрифтом и размером, что и основной текст».
    doc = _doc("Текст.", _settings(page_numbers=True))
    footer_runs = doc.sections[0].footer.paragraphs[0].runs
    assert footer_runs, "в колонтитуле должно быть поле PAGE"
    assert all(r.font.size == Pt(14) for r in footer_runs)
    assert all(r.font.name == "Times New Roman" for r in footer_runs)


# ---------- 5 «Разделы документов» ----------


def test_heading_spacing_matches_preview():
    # Заголовки БЕЗ дополнительных интервалов до/после: между заголовком и
    # текстом (и между заголовками) — обычный полуторный интервал (в превью
    # заголовки без спейсеров и padding); без разрядки символов.
    doc = _doc("# Раздел\n\n## Подраздел\n\n### Пункт\n\nТекст.")
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        assert st.paragraph_format.space_before == Pt(0)
        assert st.paragraph_format.space_after == Pt(0)
        assert st.paragraph_format.line_spacing == 1.5
        assert st.paragraph_format.keep_with_next is True
        rpr = st.element.get_or_add_rPr().xml
        assert "w:spacing" not in rpr  # разрядки символов нет


def test_heading_no_trailing_dot():
    # «В конце цифр и текста заголовка точки НЕ ставятся».
    doc = _doc("# Введение.\n\nТекст.\n\n# Обзор литературы.\n\n## Первый подраздел.\n\nТекст.")
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "ВВЕДЕНИЕ" in joined and "ВВЕДЕНИЕ." not in joined
    assert f"1{NBSP}Обзор литературы" in joined and "Обзор литературы." not in joined
    assert f"1.1{NBSP}Первый подраздел" in joined and "Первый подраздел." not in joined
    # внутренние точки не трогаем: «Приложение А. Листинг кода»
    doc2 = _doc("# Приложение А. Листинг кода\n\nТекст.")
    assert any("ПРИЛОЖЕНИЕ А. ЛИСТИНГ КОДА" == p.text for p in doc2.paragraphs)


def test_structural_headings_caps_centered_new_page():
    doc = _doc("# Введение\n\nТекст.\n\n# Заключение\n\nЕщё текст.")
    intro = next(p for p in doc.paragraphs if p.text == "ВВЕДЕНИЕ")
    concl = next(p for p in doc.paragraphs if p.text == "ЗАКЛЮЧЕНИЕ")
    for p in (intro, concl):
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert p.paragraph_format.first_line_indent == Cm(0)
    # Каждый структурный элемент — с нового листа (у первого разрыв не нужен).
    assert concl.paragraph_format.page_break_before is True


# ---------- 3 «Реферат» и 4 «Содержание» ----------


def test_referat_placed_before_toc_and_excluded_from_it():
    md = (
        "# Реферат\n\nПояснительная записка 62 с., 26 рис., 9 источников.\n\n"
        "# Введение\n\nТекст введения.\n"
    )
    doc = _doc(md, _settings(toc=True))
    texts = [p.text for p in doc.paragraphs]
    i_ref = texts.index("РЕФЕРАТ")
    i_toc = texts.index("СОДЕРЖАНИЕ")
    i_intro = texts.index("ВВЕДЕНИЕ")
    assert i_ref < i_toc < i_intro  # титул → реферат → содержание → введение
    # Реферат не попадает в содержание: уровень структуры понижен до текста.
    ref_p = doc.paragraphs[i_ref]
    assert 'w:outlineLvl' in ref_p._p.xml and 'w:val="9"' in ref_p._p.xml
    # Содержание после реферата начинается с новой страницы.
    assert doc.paragraphs[i_toc].paragraph_format.page_break_before is True


def test_toc_entry_styles_indents():
    # Подразделы в содержании сдвигаются на 0,5 см, пункты — на 1 см
    # (в превью — buildTocRow: 0/5/10 мм).
    # Сравнение в мм с допуском: отступы хранятся в twips с округлением.
    doc = _doc("# Раздел\n\nТекст.", _settings(toc=True))
    assert doc.styles["TOC 1"].paragraph_format.left_indent == Cm(0)
    assert abs(doc.styles["TOC 2"].paragraph_format.left_indent.mm - 5) < 0.05
    assert abs(doc.styles["TOC 3"].paragraph_format.left_indent.mm - 10) < 0.05
    for name in ("TOC 1", "TOC 2", "TOC 3"):
        st = doc.styles[name]
        assert st.font.name == "Times New Roman"
        assert st.font.size == Pt(14)
        assert st.paragraph_format.line_spacing == 1.5
        # Без наследования docDefaults (space_after 8 пт) — записи содержания
        # идут строка к строке, как в превью.
        assert st.paragraph_format.space_before == Pt(0)
        assert st.paragraph_format.space_after == Pt(0)


# ---------- 6 «Перечисления» и 11 «Список источников» ----------


def test_bibliography_number_with_dot_plain_indent():
    # Номер источника С ТОЧКОЙ («1.») — требование пользователя (вуз);
    # запись с абзацного отступа, как обычный текст — без висячего отступа.
    md = "# Список использованных источников\n\n1. Иванов И. И. Книга. — М., 2020.\n"
    doc = _doc(md)
    p = next(p for p in doc.paragraphs if "Иванов" in p.text)
    assert p.text.startswith(f"1.{NBSP}Иванов")
    pf = p.paragraph_format
    assert round(pf.first_line_indent.mm, 1) == 12.5
    assert pf.left_indent is None
    assert "\t" not in p.text


def test_formula_gde_each_line_with_indent():
    # Пояснения к формуле: «где …» и каждое следующее — отдельным абзацем
    # с красной строки.
    md = "$$S = a b$$\n\nгде S — площадь; a — ширина; b — высота.\n"
    doc = _doc(md)
    lines = [p for p in doc.paragraphs if "—" in p.text or "–" in p.text]
    gde = [p for p in lines if p.text]
    assert [p.text for p in gde] == ["где S – площадь;", "a – ширина;", "b – высота."]
    for p in gde:
        assert round(p.paragraph_format.first_line_indent.mm, 1) == 12.5
        assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_list_markers_dash_and_digit_paren():
    doc = _doc("- первый пункт;\n- второй пункт.\n\n1. раз;\n2. два.\n")
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert f"–{NBSP}первый пункт;" in joined
    assert f"1){NBSP}раз;" in joined and f"2){NBSP}два." in joined
    # Пункт — с красной строки как обычный абзац; продолжение длинного пункта
    # переносится к левому полю (без висячего отступа и табуляций).
    p = next(p for p in doc.paragraphs if "первый пункт" in p.text)
    pf = p.paragraph_format
    assert round(pf.first_line_indent.mm, 1) == 12.5
    assert pf.left_indent is None
    assert "\t" not in p.text


# ---------- 7 «Рисунки» ----------


def test_figure_caption_single_line_spacing():
    # Многострочное наименование рисунка — через один межстрочный интервал.
    md = "Рисунок: Схема\n![Схема](нет-такого-ассета)\n"
    doc = _doc(md)
    cap = next(p for p in doc.paragraphs if p.text.startswith("Рисунок 1"))
    assert cap.paragraph_format.line_spacing == 1.0
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_table_caption_single_spacing_tight_to_table():
    # Подпись таблицы: многострочная — через один межстрочный интервал,
    # к самой таблице — вплотную (нулевой интервал). Зеркало превью:
    # line-height:1.15 у capHtml в gostRender.ts.
    md = "Таблица: Сравнение\n| А | Б |\n|---|---|\n| 1 | 2 |\n"
    doc = _doc(md)
    cap = next(p for p in doc.paragraphs if p.text.startswith("Таблица 1"))
    pf = cap.paragraph_format
    assert pf.line_spacing == 1.0
    assert pf.space_after == Pt(0)
    assert pf.keep_with_next is True


def test_table_header_keeps_with_first_body_row():
    # Шапка не остаётся одна внизу страницы: keep_with_next в ячейках шапки
    # прилипляет её к первой строке тела (в превью — minFirst в paginate.ts).
    # Без этого LibreOffice при экспорте в PDF ТЕРЯЛ все строки тела таблицы,
    # разорванной сразу после tblHeader-строки.
    md = "Таблица: Сравнение\n| А | Б |\n|---|---|\n| 1 | 2 |\n"
    doc = _doc(md)
    table = doc.tables[0]
    for cell in table.rows[0].cells:
        assert cell.paragraphs[0].paragraph_format.keep_with_next is True
    for cell in table.rows[1].cells:
        assert cell.paragraphs[0].paragraph_format.keep_with_next is None


def test_placeholder_box_matches_preview():
    # Заглушка отсутствующей картинки — рамка 125×62 мм, как в превью
    # (иначе высота блока и раскладка страниц расходятся).
    md = "Рисунок: Схема\n![Схема](нет-такого-ассета)\n"
    doc = _doc(md)
    xml = doc.tables[0]._tbl.xml
    assert 'w:hRule="exact"' in xml
    assert f'w:val="{round(62 / 25.4 * 1440)}"' in xml  # высота 62 мм
    assert f'w:w="{round(125 / 25.4 * 1440)}"' in xml  # ширина 125 мм


def test_image_without_dpi_treated_as_96():
    # PNG без pHYs: python-docx подставляет 72 dpi, а превью считает 96 —
    # конвертер обязан считать так же (иначе картинка в DOCX крупнее на 1/3).
    tiny_png = (
        "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
        "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    doc = _doc("![Точка](asset:dot)", assets={"asset:dot": tiny_png})
    shape = doc.inline_shapes[0]
    assert abs(shape.width.mm - 25.4 / 96) < 0.01  # 1 px при 96 dpi


# ---------- Полный отчёт (фикстура) ----------


def test_kalman_report_builds_by_gost():
    doc = _doc(FIXTURE, _settings(title_page=True, toc=True, page_numbers=True))
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)

    # Структурные заголовки — прописными; разделы — нумерованные.
    for structural in ("ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"):
        assert structural in texts
    assert "ПРИЛОЖЕНИЕ А. ЛИСТИНГ КОДА" in texts
    assert f"1{NBSP}Теоретические основы оптимальной фильтрации" in texts
    assert f"1.1{NBSP}Задача оптимальной фильтрации и её постановка" in texts

    # Каждый раздел — с нового листа: 8 h1 после первого + СОДЕРЖАНИЕ.
    breaks = sum(
        1 for p in doc.paragraphs if p.paragraph_format.page_break_before
    )
    assert breaks >= 8

    # Подписи таблиц и рисунков с автонумерацией.
    assert "Таблица 1 – Сравнение методов оценивания состояния динамических систем" in joined
    assert "Рисунок 1 – Классификация методов оптимальной фильтрации" in joined

    # Формулы нумеруются по порядку, номер в круглых скобках.
    assert "(1)" in joined and "(2)" in joined

    # Источники: номер с точкой, с абзацного отступа.
    src = next(p for p in doc.paragraphs if "Калман Р. Э." in p.text)
    assert src.text.startswith(f"1.{NBSP}")
    assert round(src.paragraph_format.first_line_indent.mm, 1) == 12.5
