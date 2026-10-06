import base64
import io
import socket
import shutil

import httpx
import pytest
from docx import Document

from app.convert import images
from app.convert.gost import build_docx
from app.convert.md_parser import parse_markdown, parse_inline
from app.convert.models import GostSettings


def test_empty_edge_columns_preserve_positions():
    rows = parse_markdown('||B||\n|---|---|---|\n|1|2|3|')[0].rows
    assert rows == [['', 'B', ''], ['1', '2', '3']]
    doc = Document(io.BytesIO(build_docx('||B||\n|---|---|---|\n|1|2|3|', GostSettings(title_page=False, toc=False), {})))
    assert [cell.text for cell in doc.tables[0].rows[0].cells] == ['', 'B', '']


def test_nested_math_and_code_preserve_types():
    runs = parse_inline('**Текст $x^2$ и `code`**')
    assert next(r for r in runs if r.math).math == 'x^2'
    assert next(r for r in runs if r.math).bold
    assert next(r for r in runs if r.code).text == 'code'
    assert next(r for r in runs if r.code).bold


def _png_with_dpi(dpi=300):
    # Строим валидный PNG стандартной библиотекой, без Pillow.
    import struct
    import zlib
    def chunk(kind, data):
        return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!IIBBBBB', 300, 300, 8, 2, 0, 0, 0)) + chunk(b'pHYs', struct.pack('!IIB', round(dpi / .0254), round(dpi / .0254), 1)) + chunk(b'IDAT', zlib.compress((b'\0' + b'\xff' * 900) * 300)) + chunk(b'IEND', b'')


def test_regular_image_uses_css_pixels_mermaid_keeps_raster_scale():
    assets = {'asset:test': base64.b64encode(_png_with_dpi()).decode()}
    settings = GostSettings(title_page=False, toc=False)
    doc = Document(io.BytesIO(build_docx('![test](asset:test)', settings, assets)))
    assert doc.inline_shapes[0].width.mm == pytest.approx(79.375, abs=.01)
    doc = Document(io.BytesIO(build_docx('```mermaid\nflowchart LR\nA-->B\n```', settings, {'mermaid-0': assets['asset:test']})))
    assert doc.inline_shapes[0].width.mm == pytest.approx(25.4, abs=.01)


def _mock_network(monkeypatch, handler):
    original = httpx.Client
    monkeypatch.setattr(images.socket, 'getaddrinfo', lambda *a, **k: [(socket.AF_INET, socket.SOCK_STREAM, 6, '', ('93.184.216.34', 443))])
    monkeypatch.setattr(images.httpx, 'Client', lambda **kw: original(transport=httpx.MockTransport(handler), **kw))


def test_redirect_cannot_reach_private_host(monkeypatch):
    calls = []
    def handler(req):
        calls.append(req)
        return httpx.Response(302, headers={'location': 'http://127.0.0.1/internal'})
    _mock_network(monkeypatch, handler)
    assert images._fetch_url('https://public.example/image') is None
    assert len(calls) == 1
    assert calls[0].url.host == '93.184.216.34'
    assert calls[0].headers['Host'] == 'public.example'
    assert calls[0].extensions['sni_hostname'] == 'public.example'


@pytest.mark.parametrize('ip', ['127.0.0.1', '10.1.2.3', '169.254.169.254', '::1', '224.0.0.1'])
def test_private_and_multicast_addresses_rejected(monkeypatch, ip):
    monkeypatch.setattr(images.socket, 'getaddrinfo', lambda *a, **k: [(socket.AF_INET, socket.SOCK_STREAM, 6, '', (ip, 443))])
    assert images._public_address('example.test', 443) is None


def test_image_stream_stops_at_limit(monkeypatch):
    emitted = []
    class Stream(httpx.SyncByteStream):
        def __iter__(self):
            for _ in range(10):
                emitted.append(1)
                yield b'x' * 65536
    _mock_network(monkeypatch, lambda req: httpx.Response(200, headers={'content-type': 'image/png'}, stream=Stream()))
    monkeypatch.setattr(images, 'MAX_IMAGE_BYTES', 65536)
    assert images._fetch_url('https://public.example/image') is None
    assert len(emitted) == 2


@pytest.mark.skipif(shutil.which('pandoc') is None, reason='Pandoc проверяется обязательно в Docker/CI')
def test_real_pandoc_produces_native_math():
    doc = Document(io.BytesIO(build_docx('**$x^2$**\n\n$$E=mc^2$$', GostSettings(title_page=False, toc=False), {})))
    assert len(doc.element.xpath('//m:oMath')) == 2


def test_encoded_responses_rejected_before_decompression(monkeypatch):
    _mock_network(monkeypatch, lambda req: httpx.Response(200, headers={'content-type': 'image/png', 'content-encoding': 'gzip'}, content=b''))
    assert images._fetch_url('https://public.example/image') is None


@pytest.mark.parametrize('text', ['Текст\x00', 'Текст\x0b', 'Текст\ud800', 'Текст\uffff'])
def test_xml_invalid_text_rejected_before_conversion(text):
    from pydantic import ValidationError
    from app.convert.models import ConvertRequest
    with pytest.raises(ValidationError, match='управляющие символы'):
        ConvertRequest(markdown=text)
    with pytest.raises(ValidationError, match='управляющие символы'):
        GostSettings(topic=text)


def test_document_preserves_unicode_tabs_and_newlines():
    from app.convert.models import ConvertRequest
    req = ConvertRequest(markdown='Текст\tс символом 😀\n\nЕщё абзац', settings=GostSettings(title_page=False, toc=False))
    doc = Document(io.BytesIO(build_docx(req.markdown, req.settings, {})))
    assert '😀' in '\n'.join(p.text for p in doc.paragraphs)
