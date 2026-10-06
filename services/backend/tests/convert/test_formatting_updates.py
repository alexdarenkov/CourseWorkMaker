"""Парные проверки согласованных правил: frontend/tests/lib/formattingUpdates.test.ts."""
import io

import pytest
from docx import Document
from docx.shared import Pt

from app.convert.gost import build_docx
from app.convert.md_parser import parse_markdown, ListBlock
from app.convert.models import GostSettings


def doc(md, auto_number=True):
    return Document(io.BytesIO(build_docx(md, GostSettings(title_page=False, toc=False, page_numbers=False, auto_number=auto_number), {})))


@pytest.mark.parametrize('cols', [3, 4, 6, 7])
def test_table_font_and_header(cols):
    row = '| ' + ' | '.join(['Текст'] * cols) + ' |'
    table = doc(row + '\n|' + '---|' * cols + '\n' + row).tables[0]
    for index, row in enumerate(table.rows):
        run = row.cells[0].paragraphs[0].runs[0]
        assert run.font.size == Pt(14 if cols < 4 else 12)
        assert bool(run.bold) == (index == 0)


def test_inline_code_font_and_content():
    run = doc('` a  b `').paragraphs[0].runs[0]
    assert run.font.size == Pt(12)
    assert run.font.name == 'Courier New'
    assert run.text == ' a  b '


def test_unnumbered_table_caption():
    document = doc('Таблица: Опыт\n| А |\n|---|\n| Б |', False)
    texts = [p.text for p in document.paragraphs]
    assert 'Опыт' in texts
    assert not any('Таблица 1' in t for t in texts)


def test_empty_figure_caption():
    texts = [p.text for p in doc('![](missing)').paragraphs]
    assert 'Рисунок 1' in texts
    assert not any('Рисунок 1 -' in t for t in texts)


def test_alpha_list_after_text():
    md = 'Текст\nа) Первый;\nб) Второй.\n\n1) Новый.'
    blocks = parse_markdown(md)
    assert isinstance(blocks[1], ListBlock)
    assert blocks[1].markers == ['а)', 'б)']
    texts = [p.text for p in doc(md).paragraphs]
    assert texts == ['Текст', 'а)\u00a0Первый;', 'б)\u00a0Второй.', '1)\u00a0Новый.']


@pytest.mark.parametrize('md', ['$$x=1$$', '![](missing)', '```\nx\n```', '| А |\n|---|\n| Б |'])
def test_trailing_free_line(md):
    assert doc(md).paragraphs[-1].text == ''


def test_math_repeats_binary_operator_at_wrap():
    settings = doc('$$a+b=c$$').settings.element
    assert settings.xpath('./m:mathPr/m:brkBin/@m:val') == ['repeat']
    assert settings.xpath('./m:mathPr/m:brkBinSub/@m:val') == ['--']
