import io

from docx import Document

from app.convert.gost import build_docx
from app.convert.md_parser import (
    Blank,
    CodeBlock,
    Figure,
    Heading,
    ListBlock,
    MathBlock,
    MermaidBlock,
    PageBreak,
    Paragraph,
    Table,
    parse_inline,
    parse_markdown,
)
from app.convert.models import GostSettings

SAMPLE = """# Введение

Текст **жирный** и *курсив* с формулой $x^2$.

# Анализ предметной области

## Обзор решений

Таблица: Сравнение
| А | Б |
|---|---|
| 1 | 2 |

Рисунок: Схема
```mermaid
flowchart LR
  A --> B
```

$$E = mc^2$$

```python
print("hi")
```

- пункт один;
- пункт два.

# Список использованных источников

1. ГОСТ 7.32-2017.
2. Второй источник.
"""


def test_parse_blocks():
    blocks = parse_markdown(SAMPLE)
    types = [type(b) for b in blocks]
    assert types == [
        Heading, Paragraph, Heading, Heading, Table, MermaidBlock,
        MathBlock, CodeBlock, ListBlock, Heading, ListBlock,
    ]
    table = blocks[4]
    assert table.caption == "Сравнение"
    assert table.rows == [["А", "Б"], ["1", "2"]]
    mermaid = blocks[5]
    assert mermaid.caption == "Схема"


def test_parse_pagebreak_and_blank():
    # «---» — разрыв страницы; лишние пустые строки (сверх одной) — видимые
    # пустые строки (N-1 штук).
    blocks = parse_markdown("Первый\n\n\n\nВторой\n\n---\n\nТретий")
    types = [type(b) for b in blocks]
    assert types == [Paragraph, Blank, Blank, Paragraph, PageBreak, Paragraph]


def test_build_docx_pagebreak():
    data = build_docx("Один\n\n---\n\nДва", GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    xml = "\n".join(p._p.xml for p in doc.paragraphs)
    assert 'w:type="page"' in xml  # принудительный разрыв страницы присутствует


def test_ordered_list_continuous_numbering():
    # Пункты с одинаковым маркером «1)» и пустыми строками между ними образуют
    # отдельные списки — нумерация должна быть сквозной (1,2,3), а не «1,1,1».
    md = "1) Первый\n\n1) Второй\n\n1) Третий\n"
    data = build_docx(md, GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "1) Первый" in joined
    assert "2) Второй" in joined
    assert "3) Третий" in joined


def test_bibliography_continuous_numbering():
    # Источники нумеруются арабскими цифрами С ТОЧКОЙ (требование вуза;
    # выписка ГОСТ 6.16 говорит «без точки»), сквозной нумерацией через
    # пустые строки между пунктами.
    md = (
        "# Список использованных источников\n\n"
        "1. Первый источник.\n\n1. Второй источник.\n\n1. Третий источник.\n"
    )
    data = build_docx(md, GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "1. Первый источник." in joined
    assert "2. Второй источник." in joined
    assert "3. Третий источник." in joined


def test_nested_inline_no_control_chars():
    # Вложенное форматирование (код/формула внутри **жирного**) не должно
    # оставлять служебный \x00 в тексте ранов — иначе DOCX не собирается.
    for src in ("**жирный с `кодом` внутри**", "**текст $x^2$ формула**", "*курсив с `code`*"):
        runs = parse_inline(src)
        joined = "".join((r.text or "") + (r.math or "") for r in runs)
        assert "\x00" not in joined
    data = build_docx(
        "**жирный с `кодом` внутри** и $a+b$",
        GostSettings(title_page=False, toc=False),
        {},
    )
    doc = Document(io.BytesIO(data))
    assert "кодом" in "\n".join(p.text for p in doc.paragraphs)


def test_parse_inline():
    runs = parse_inline("a **b** *c* `d` $e$ [f](http://x)")
    flags = [(r.text, r.bold, r.italic, r.code, r.underline, r.math) for r in runs]
    assert ("b", True, False, False, False, None) in flags
    assert ("c", False, True, False, False, None) in flags
    assert ("d", False, False, True, False, None) in flags
    assert ("f", False, False, False, True, None) in flags
    assert any(r.math == "e" for r in runs)


def test_build_docx_smoke():
    settings = GostSettings(
        title_header="ТЕСТОВЫЙ УНИВЕРСИТЕТ\nКафедра тестов",
        title_work="КУРСОВАЯ РАБОТА\nпо дисциплине «Тестирование»\nна тему:",
        topic="Тема работы",
        title_people="Выполнил: студент группы Т-01\nИванов Иван\n\nРуководитель: Петров П. П.",
        title_bottom="Москва\n2026",
    )
    data = build_docx(SAMPLE, settings, {})
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)
    assert "КУРСОВАЯ РАБОТА" in joined
    assert "«Тема работы»" in joined  # тема в кавычках полужирно
    assert "Выполнил:\tстудент группы Т-01" in joined  # метка + правый таб
    assert "СОДЕРЖАНИЕ" in joined
    assert "ВВЕДЕНИЕ" in joined
    assert "1\u00a0Анализ предметной области" in joined
    assert "1.1 Обзор решений" in joined
    assert "Таблица 1 – Сравнение" in joined  # среднее тире (–), не длинное
    assert "Рисунок 1 – Схема" in joined  # mermaid без ассета -> плейсхолдер + подпись
    assert "1. ГОСТ 7.32-2017." in joined  # источник: номер с точкой (вуз)
    assert "– пункт один;" in joined  # маркер + NBSP, с красной строки
    assert len(doc.tables) == 3  # таблица данных + заглушка mermaid + листинг-рамка

    sec = doc.sections[0]
    assert round(sec.left_margin.mm) == 30
    assert round(sec.right_margin.mm) == 15
    assert round(sec.top_margin.mm) == 20
    assert round(sec.bottom_margin.mm) == 20


def test_captions_without_auto_number():
    # При выключенной автонумерации подписи таблиц и рисунков остаются без
    # «Таблица N» / «Рисунок N» — только сам текст подписи.
    md = "Таблица: Сравнение\n| А |\n|---|\n| 1 |\n\nРисунок: Схема\n![Схема](placeholder)\n"
    data = build_docx(md, GostSettings(title_page=False, toc=False, auto_number=False), {})
    doc = Document(io.BytesIO(data))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "Сравнение" in joined and "Таблица" not in joined
    assert "Схема" in joined and "Рисунок" not in joined


def test_table_fits_page_and_repeats_header():
    # Таблица занимает 100% ширины контентной области (w:tblW pct=5000) и
    # повторяет шапку на каждой странице (w:tblHeader на первой строке).
    md = "Таблица: Т\n| А | Б | В |\n|---|---|---|\n| 1 | 2 | 3 |\n| 4 | 5 | 6 |\n"
    data = build_docx(md, GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    table = doc.tables[0]
    xml = table._tbl.xml
    assert 'w:type="pct"' in xml and 'w:w="5000"' in xml
    assert xml.count("tblHeader") == 1  # только первая строка — шапка
    assert "cantSplit" in xml


def test_wide_table_uses_smaller_font():
    # >6 колонок — кегль 12 пт (ГОСТ допускает); узкая таблица — обычные 14 пт.
    def table_md(cols):
        header = "| " + " | ".join(f"К{i}" for i in range(cols)) + " |"
        sep = "|" + "---|" * cols
        row = "| " + " | ".join(str(i) for i in range(cols)) + " |"
        return f"Таблица: Т\n{header}\n{sep}\n{row}\n"

    from docx.shared import Pt

    for cols, expected in ((3, Pt(14)), (8, Pt(12))):
        data = build_docx(table_md(cols), GostSettings(title_page=False, toc=False), {})
        doc = Document(io.BytesIO(data))
        run = doc.tables[0].cell(0, 0).paragraphs[0].runs[0]
        assert run.font.size == expected, f"{cols} колонок → ждали {expected}"


def test_list_red_line_no_hanging_indent():
    # Пункт перечисления — как обычный абзац с красной строки: маркер на
    # INDENT (1.25см), продолжение длинного пункта переносится к ЛЕВОМУ полю
    # (left_indent нет — висячий отступ убран). Сравнение в мм с округлением —
    # EMU/twips дают долю мм погрешности.
    md = "- пункт;\n1. первый;\n"
    data = build_docx(md, GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    dash_p = next(p for p in doc.paragraphs if p.text.startswith("–"))
    num_p = next(p for p in doc.paragraphs if p.text.startswith("1)"))
    for p in (dash_p, num_p):
        assert abs(p.paragraph_format.first_line_indent.mm - 12.5) < 0.1
        assert p.paragraph_format.left_indent is None
        assert "\t" not in p.text


def test_adjacent_formulas_no_free_lines_between():
    # Формулы подряд идут вплотную; свободные строки — только вокруг группы
    # (перед первой формулой и после последней, перед текстом). Свободная
    # строка — НАСТОЯЩИЙ пустой абзац (перенос строки), а не межабзацный
    # интервал: нормоконтроль проверяет именно пустые строки.
    from docx.shared import Pt

    md = "Текст до.\n\n$$a = 1$$\n\n$$b = 2$$\n\nТекст после.\n"
    data = build_docx(md, GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    i_before = texts.index("Текст до.")
    i_f1 = next(i for i, t in enumerate(texts) if "(1)" in t)
    i_f2 = next(i for i, t in enumerate(texts) if "(2)" in t)
    i_after = texts.index("Текст после.")
    # Между текстом и первой формулой — ровно один пустой абзац.
    assert [t.strip() for t in texts[i_before + 1 : i_f1]] == [""]
    # Между формулами подряд пустых абзацев нет.
    assert i_f2 == i_f1 + 1
    # После последней формулы перед текстом — ровно один пустой абзац.
    assert [t.strip() for t in texts[i_f2 + 1 : i_after]] == [""]
    # Формулы без межабзацных интервалов (свободные строки — абзацами).
    for f in (doc.paragraphs[i_f1], doc.paragraphs[i_f2]):
        assert not f.paragraph_format.space_before or f.paragraph_format.space_before == Pt(0)


def test_formula_keeps_with_gde():
    # Формула с пояснением «где …» не отрывается от него при переносе страницы.
    md = "$$E = mc^2$$\n\nгде E — энергия; m — масса.\n"
    data = build_docx(md, GostSettings(title_page=False, toc=False), {})
    doc = Document(io.BytesIO(data))
    formula_p = next(p for p in doc.paragraphs if "(1)" in p.text)
    assert formula_p.paragraph_format.keep_with_next is True
    # Формула без «где» — без keep_with_next.
    data2 = build_docx("$$y = x$$\n\nОбычный абзац.\n", GostSettings(title_page=False, toc=False), {})
    doc2 = Document(io.BytesIO(data2))
    formula_p2 = next(p for p in doc2.paragraphs if "(1)" in p.text)
    assert formula_p2.paragraph_format.keep_with_next is None


def test_build_docx_no_extras():
    settings = GostSettings(
        title_page=False, toc=False, page_numbers=False,
        bibliography=False, auto_number=False,
    )
    data = build_docx("# Введение\n\nПросто текст.", settings, {})
    doc = Document(io.BytesIO(data))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "КУРСОВАЯ РАБОТА" not in joined
    assert "СОДЕРЖАНИЕ" not in joined
    assert "ВВЕДЕНИЕ" in joined
