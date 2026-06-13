"""Извлечение текста из загруженных файлов-источников."""

from __future__ import annotations

import io
import logging

from . import config

log = logging.getLogger(__name__)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            return _from_pdf(data)
        if name.endswith(".docx"):
            return _from_docx(data)
        return data.decode("utf-8", errors="replace")
    except Exception:
        log.warning("Failed to extract text from %s", filename, exc_info=True)
        return ""


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def build_context(files: list[tuple[str, bytes]]) -> str:
    """Собирает единый текстовый контекст из файлов с ограничением размера."""
    chunks: list[str] = []
    budget = config.MAX_CONTEXT_CHARS
    per_file = budget // max(len(files), 1)
    for filename, data in files:
        text = extract_text(filename, data).strip()
        if not text:
            continue
        if len(text) > per_file:
            text = text[:per_file] + "\n…[обрезано]"
        chunks.append(f"===== Файл: {filename} =====\n{text}")
    return "\n\n".join(chunks)[:budget]
